"""Generate AI-CoE-Horizon-Benchmark.xlsx.

Companion workbook for issue #14 — comparator on top of Horizon Assessment.
Inputs six pillar scores; outputs percentile per pillar against v1 RSA cohort (N=20).
"""

from __future__ import annotations
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent / "AI-CoE-Horizon-Benchmark.xlsx"

BLUE = "0067B8"
NAVY = "0B1F3A"
LIGHT = "EAF4FB"
WHITE = "FFFFFF"

THIN = Side(border_style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PILLARS = [
    # (code, name, min, p25, p50, p75, max, mean)
    ("P1", "Business strategy",          1.0, 2.0, 2.5, 3.5, 4.5, 2.7),
    ("P2", "Org & culture",              1.0, 1.5, 2.0, 2.5, 4.0, 2.1),
    ("P3", "AI strategy & experience",   1.0, 2.0, 2.5, 3.0, 4.0, 2.5),
    ("P4", "Tech & data",                1.5, 2.0, 3.0, 3.5, 4.5, 2.9),
    ("P5", "Governance & security",      1.0, 1.5, 2.5, 3.0, 4.0, 2.4),
    ("P+1", "Co-sell & partner",         1.0, 1.5, 2.0, 3.0, 4.0, 2.2),
]


def header(cell, fill=BLUE, color=WHITE, size=11):
    cell.font = Font(name="Segoe UI", bold=True, color=color, size=size)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = BORDER


def style(cell, bold=False, size=10, fill=None, align="left"):
    cell.font = Font(name="Segoe UI", bold=bold, size=size, color="000000")
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal=align, vertical="top", wrap_text=True)
    cell.border = BORDER


def widths(ws, cols):
    for i, w in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def sheet_readme(wb):
    ws = wb.create_sheet("README")
    ws.sheet_view.showGridLines = False
    widths(ws, [110])
    rows = [
        ("AI CoE — Horizon Benchmark (v1)", True, 18, NAVY),
        ("Yuri Baijnath — CSU Cloud & AI Lead (South Africa), Microsoft · v1.0 · 2026-06-03", False, 11, NAVY),
        ("", False, 11, NAVY),
        ("Companion to: AI-CoE-Horizon-Assessment.xlsx and ai-coe-horizon-benchmark-dataset.md", False, 11, NAVY),
        ("Closes issue #14 (benchmark dataset for Horizon Assessment — comparative scoring).", False, 11, NAVY),
        ("", False, 11, NAVY),
        ("Workflow", True, 14, BLUE),
        ("1. Score the customer in AI-CoE-Horizon-Assessment.xlsx (six pillars, 1-5 scale).", False, 11, NAVY),
        ("2. Enter the six pillar scores in the 'Input' sheet of this workbook.", False, 11, NAVY),
        ("3. The 'Comparator' sheet computes percentile per pillar against the v1 RSA cohort (N=20).", False, 11, NAVY),
        ("4. The 'Output' sheet is a one-page summary — exportable as PDF for the steering pack.", False, 11, NAVY),
        ("5. The 'Benchmark' sheet holds the v1 quartile points; do not edit without CoE lead sign-off.", False, 11, NAVY),
        ("", False, 11, NAVY),
        ("Methodology notes", True, 14, BLUE),
        ("Sample: 20 RSA organisations · BFSI 6 · Telco 3 · Public sector / SOE 4 · Retail 3 · Mining 2 · Healthcare 2.", False, 11, NAVY),
        ("Method: piecewise-linear interpolation between Min / P25 / P50 / P75 / Max anchor points.", False, 11, NAVY),
        ("Stability: v1 N=20 supports overall percentiles only — NOT sector × size cuts (need N>=50).", False, 11, NAVY),
        ("Refresh: quarterly; trend annually.", False, 11, NAVY),
        ("", False, 11, NAVY),
        ("Privacy", True, 14, BLUE),
        ("Customers consent via Horizon Assessment engagement letter. Cells with N<5 never published.", False, 11, NAVY),
        ("Raw per-organisation scores live in Dataverse, accessible to CoE lead + assessors only.", False, 11, NAVY),
    ]
    for i, (text, bold, size, color) in enumerate(rows, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(name="Segoe UI", bold=bold, size=size, color=color)
        c.alignment = Alignment(wrap_text=True, vertical="top")


def sheet_benchmark(wb):
    ws = wb.create_sheet("Benchmark")
    ws.sheet_view.showGridLines = False
    widths(ws, [8, 32, 10, 10, 10, 10, 10, 10])
    hdrs = ["Code", "Pillar", "Min", "P25", "P50", "P75", "Max", "Mean"]
    for i, h in enumerate(hdrs, start=1):
        header(ws.cell(row=1, column=i, value=h))
    for r, (code, name, mn, p25, p50, p75, mx, mean) in enumerate(PILLARS, start=2):
        style(ws.cell(row=r, column=1, value=code), bold=True, fill=LIGHT)
        style(ws.cell(row=r, column=2, value=name), fill=LIGHT)
        for c, v in enumerate([mn, p25, p50, p75, mx, mean], start=3):
            style(ws.cell(row=r, column=c, value=v), align="center")


def sheet_input(wb):
    ws = wb.create_sheet("Input")
    ws.sheet_view.showGridLines = False
    widths(ws, [8, 32, 16])
    for i, h in enumerate(["Code", "Pillar", "Customer score (1-5)"], start=1):
        header(ws.cell(row=1, column=i, value=h))
    for r, (code, name, *_rest) in enumerate(PILLARS, start=2):
        style(ws.cell(row=r, column=1, value=code), bold=True, fill=LIGHT)
        style(ws.cell(row=r, column=2, value=name), fill=LIGHT)
        c = ws.cell(row=r, column=3, value=2.5)
        style(c, bold=True, align="center", fill="FFF2CC")


def sheet_comparator(wb):
    """Piecewise-linear percentile.

    Anchors: Min=0, P25=25, P50=50, P75=75, Max=100.
    score <= Min  -> 0
    score >= Max  -> 100
    Min < score <= P25 -> linear between 0 and 25
    P25 < score <= P50 -> linear between 25 and 50
    ... etc.
    """
    ws = wb.create_sheet("Comparator")
    ws.sheet_view.showGridLines = False
    widths(ws, [8, 28, 14, 10, 10, 10, 10, 10, 14, 16])
    hdrs = ["Code", "Pillar", "Customer score", "Min", "P25", "P50", "P75", "Max", "Percentile", "Position"]
    for i, h in enumerate(hdrs, start=1):
        header(ws.cell(row=1, column=i, value=h))

    for r in range(2, 8):
        style(ws.cell(row=r, column=1, value=f"=Input!A{r}"), bold=True, fill=LIGHT)
        style(ws.cell(row=r, column=2, value=f"=Input!B{r}"), fill=LIGHT)
        style(ws.cell(row=r, column=3, value=f"=Input!C{r}"), bold=True, align="center")
        for c, src in enumerate(["C", "D", "E", "F", "G"], start=4):
            style(ws.cell(row=r, column=c, value=f"=Benchmark!{src}{r}"), align="center")
        s = f"C{r}"  # customer score
        mn, p25, p50, p75, mx = f"D{r}", f"E{r}", f"F{r}", f"G{r}", f"H{r}"
        pct = (
            f'=IF({s}<={mn},0,'
            f'IF({s}<={p25},25*({s}-{mn})/MAX(0.0001,({p25}-{mn})),'
            f'IF({s}<={p50},25+25*({s}-{p25})/MAX(0.0001,({p50}-{p25})),'
            f'IF({s}<={p75},50+25*({s}-{p50})/MAX(0.0001,({p75}-{p50})),'
            f'IF({s}<={mx},75+25*({s}-{p75})/MAX(0.0001,({mx}-{p75})),100)))))'
        )
        cell = ws.cell(row=r, column=9, value=pct)
        style(cell, bold=True, align="center")
        cell.number_format = "0"
        pos = f'=IF(I{r}>=75,"Leading",IF(I{r}>=50,"Above median",IF(I{r}>=25,"Below median","Lagging")))'
        style(ws.cell(row=r, column=10, value=pos), bold=True, align="center")


def sheet_output(wb):
    ws = wb.create_sheet("Output")
    ws.sheet_view.showGridLines = False
    widths(ws, [8, 32, 16, 16, 18])
    ws.cell(row=1, column=1, value="Your AI Horizon vs RSA peers").font = Font(name="Segoe UI", bold=True, size=18, color=NAVY)
    ws.cell(row=2, column=1, value="One-page summary · v1 cohort N=20 RSA organisations · 2026-06-03").font = Font(name="Segoe UI", italic=True, size=10, color=NAVY)

    for i, h in enumerate(["Code", "Pillar", "Customer", "Percentile", "Position"], start=1):
        header(ws.cell(row=4, column=i, value=h))
    for r in range(5, 11):
        src = r - 3  # Comparator rows 2..7
        style(ws.cell(row=r, column=1, value=f"=Comparator!A{src}"), bold=True, fill=LIGHT)
        style(ws.cell(row=r, column=2, value=f"=Comparator!B{src}"), fill=LIGHT)
        style(ws.cell(row=r, column=3, value=f"=Comparator!C{src}"), bold=True, align="center")
        c = ws.cell(row=r, column=4, value=f"=Comparator!I{src}")
        style(c, bold=True, align="center")
        c.number_format = "0"
        style(ws.cell(row=r, column=5, value=f"=Comparator!J{src}"), bold=True, align="center")

    ws.cell(row=12, column=1, value="Leading dimensions (P75+):").font = Font(name="Segoe UI", bold=True, size=12, color=BLUE)
    ws.cell(row=12, column=2, value='=TEXTJOIN(", ",TRUE,IF(Comparator!I2:I7>=75,Comparator!B2:B7,""))').font = Font(name="Segoe UI", size=11)
    ws.cell(row=13, column=1, value="Lagging dimensions (P25-):").font = Font(name="Segoe UI", bold=True, size=12, color=BLUE)
    ws.cell(row=13, column=2, value='=TEXTJOIN(", ",TRUE,IF(Comparator!I2:I7<25,Comparator!B2:B7,""))').font = Font(name="Segoe UI", size=11)
    ws.cell(row=15, column=1, value="Recommended starting point:").font = Font(name="Segoe UI", bold=True, size=12, color=BLUE)
    ws.cell(row=15, column=2, value="Start where you are lagging your peers, not where the brochure says.").font = Font(name="Segoe UI", italic=True, size=11)


def main():
    wb = Workbook()
    wb.remove(wb.active)
    sheet_readme(wb)
    sheet_benchmark(wb)
    sheet_input(wb)
    sheet_comparator(wb)
    sheet_output(wb)
    wb.save(str(OUT))
    print(f"Wrote {OUT.name}")


if __name__ == "__main__":
    main()
