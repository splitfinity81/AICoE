"""Build the AI CoE Objection Handling Deck.

Generates AI-CoE-Objection-Handling.pptx from
ai-coe-objection-handling-report.md. Each pillar (5+1) gets an overview
slide and a tabular objection menu listing recommended responses and the
proof point a seller should reach for in the existing AI CoE artefact pack.

Author: Yuri Baijnath — CSU Cloud & AI Lead (South Africa), Microsoft
"""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

# ---------------------------------------------------------------------------
# Theme - aligned to AI-CoE-Pitch-Deck style canon (matches build_vbd_deck.py)
# ---------------------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

C_BG       = RGBColor(0xFF, 0xFF, 0xFF)
C_INK      = RGBColor(0x1A, 0x1A, 0x1A)
C_MUTED    = RGBColor(0x55, 0x55, 0x55)
C_RULE     = RGBColor(0xD0, 0xD0, 0xD0)
C_ACCENT   = RGBColor(0x00, 0x67, 0xB8)   # Microsoft blue
C_ACCENT2  = RGBColor(0x50, 0xE6, 0xFF)   # Azure cyan
C_HIGHLIGHT= RGBColor(0x74, 0x2A, 0x9B)   # Copilot purple
C_GOOD     = RGBColor(0x10, 0x7C, 0x10)
C_WARN     = RGBColor(0xB7, 0x6E, 0x00)
C_BAND     = RGBColor(0xF3, 0xF6, 0xFA)
C_BAND2    = RGBColor(0xE8, 0xEE, 0xF7)
C_DARK     = RGBColor(0x0B, 0x1F, 0x3A)

TOTAL_SLIDES = 21

# Pillar colour assignments (used across overview cards + table headers)
PILLAR_COLORS = {
    1: C_ACCENT,
    2: C_HIGHLIGHT,
    3: C_DARK,
    4: C_GOOD,
    5: C_WARN,
    6: C_ACCENT2,   # +1 Co-sell & Partner
}

# ---------------------------------------------------------------------------
# Primitives (mirrors build_vbd_deck.py helpers)
# ---------------------------------------------------------------------------

def add_rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    shape.shadow.inherit = False
    return shape


def add_text(slide, x, y, w, h, text, *, size=14, bold=False, color=C_INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False,
             font="Segoe UI"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = Inches(0.0)
    tf.margin_right = Inches(0.0)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return tb


def add_bullets(slide, x, y, w, h, items, *, size=14, color=C_INK,
                bullet_color=C_ACCENT, line_spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.0)
    tf.margin_top = Inches(0.0)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        r1 = p.add_run()
        r1.text = "\u2022  "
        r1.font.name = "Segoe UI"
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = bullet_color
        r2 = p.add_run()
        r2.text = item
        r2.font.name = "Segoe UI"
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
    return tb


def add_section_tag(slide, text, color=C_ACCENT):
    add_text(slide, Inches(0.6), Inches(0.35), Inches(10), Inches(0.3),
             text.upper(), size=11, bold=True, color=color)


def add_title(slide, title, subtitle=None):
    add_text(slide, Inches(0.6), Inches(0.65), Inches(12.2), Inches(0.8),
             title, size=28, bold=True, color=C_INK)
    if subtitle:
        add_text(slide, Inches(0.6), Inches(1.4), Inches(12.2), Inches(0.55),
                 subtitle, size=14, color=C_MUTED)


def add_footer(slide, slide_no):
    add_rect(slide, Inches(0.6), Inches(0.28), Inches(12.13), Emu(9000), C_RULE)
    add_rect(slide, Inches(0.6), Inches(7.05), Inches(12.13), Emu(9000), C_RULE)
    add_text(slide, Inches(0.6), Inches(7.12), Inches(8), Inches(0.3),
             "Microsoft \u00b7 Confidential \u00b7 \u00a9 2026 Microsoft Corporation",
             size=9, color=C_MUTED)
    add_text(slide, Inches(10.7), Inches(7.12), Inches(2.0), Inches(0.3),
             f"{slide_no} / {TOTAL_SLIDES}",
             size=9, color=C_MUTED, align=PP_ALIGN.RIGHT)


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


def add_table(slide, x, y, w, h, headers, rows, *,
              header_fill=C_DARK, header_color=C_BG,
              header_size=11, body_size=10,
              col_widths=None, row_height=None,
              first_col_bold=True, alt_band=True):
    n_cols = len(headers)
    n_rows = len(rows) + 1
    tbl_shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    tbl = tbl_shape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = int(w * cw / total)
    if row_height:
        for r in tbl.rows:
            r.height = row_height
    for j, htext in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.04)
        cell.margin_bottom = Inches(0.04)
        tf = cell.text_frame
        tf.word_wrap = True
        tf.clear()
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = htext
        r.font.name = "Segoe UI"
        r.font.size = Pt(header_size)
        r.font.bold = True
        r.font.color.rgb = header_color
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_BAND if (alt_band and i % 2 == 0) else C_BG
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            tf = cell.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.name = "Segoe UI"
            r.font.size = Pt(body_size)
            r.font.color.rgb = C_INK
            r.font.bold = first_col_bold and (j == 0)
    return tbl


# ---------------------------------------------------------------------------
# Pillar content - sourced from ai-coe-objection-handling-report.md
# Each objection: (short label, persona, one-line response, proof point)
# ---------------------------------------------------------------------------

PILLARS = [
    {
        "num": 1,
        "name": "Business Strategy",
        "subtitle": "Value, timing, unit economics \u2014 where most \"no\" lives",
        "narrative": (
            "Anchor every response in the $1 of ECIF \u2192 $3\u20135 of consumption "
            "ratio and the ATO up-to-$1M envelope that removes customer-side "
            "spend friction. Lead with the value case, not the SKU."
        ),
        "owner": "Pillar 1 lead \u00b7 escalate to CSU GM, ATU lead",
        "objections": [
            ("\"ROI on AI is unproven.\"", "CFO, CEO",
             "CoE forces a value case per use-case before build; ATO absorbs first $1M of envelope; Stage-4 (D3) reports CFO-grade outcomes quarterly.",
             "Pitch Deck ATO slide \u00b7 VBD D3"),
            ("\"We'll wait for tech to mature.\"", "CEO, CIO",
             "88% of orgs already in (McKinsey 2025); T1 envisioning is free \u2014 stop after if value case doesn't hold.",
             "Horizon Assessment \u00b7 ATO C1"),
            ("\"Copilot per-seat doesn't pencil.\"", "CFO, CIO",
             "Per-seat \u00d7 headcount is the wrong denominator; A2 produces hours-recovered-per-persona; license where 3\u00d7 cost clears.",
             "VBDs A2, A3, A4"),
            ("\"We'll build it ourselves \u2014 cheaper.\"", "CIO, CTO",
             "Building is cheap; operating safely at scale is the stall point. Co-build (B2/C18) keeps IP yours, run-time discipline ours.",
             "VBDs C18, C13 \u00b7 F1 Factory"),
            ("\"Show me an in-industry peer.\"", "CEO, BDM",
             "Lead with NTT DATA / Capgemini for SI motions; Eskom pattern for regulated. If no in-segment ref, offer NDA peer call from adjacent segment.",
             "Pitch Deck proof points \u00b7 Eskom briefing"),
            ("\"What's the exit cost?\"", "CFO, Procurement",
             "Three optionality layers: T1 envisioning is ECIF-funded; Foundry is multi-model; data stays in your tenant. Exit at any tier returns customer-owned assets.",
             "ATO C1 \u00b7 Pillar 4 Foundry story"),
        ],
    },
    {
        "num": 2,
        "name": "Organisation & Culture",
        "subtitle": "People-readiness, change fatigue, adoption stalls",
        "narrative": (
            "Only ~1 in 3 employees has been trained on the AI tools they "
            "already have (BCG 2024). The Adoption Accelerator (A3) and "
            "Skilling Navigator (D2) exist precisely for this. Measure "
            "Active Use, Sentiment, Habit \u2014 if they don't move, re-scope."
        ),
        "owner": "Pillar 2 lead \u00b7 escalate with customer HR sponsor",
        "objections": [
            ("\"Our people aren't ready.\"", "CHRO, COO, CIO",
             "Modal state \u2014 BCG 2024 confirms. A3 runs 8\u201312 weeks with champions + measured uplift. Don't roll wide until metrics move.",
             "VBD A3 \u00b7 VBD D2"),
            ("\"We have change fatigue from [other rollout].\"", "COO, LOB head",
             "Sequence accordingly \u2014 thin first wave, one persona, fold into existing change programme. Co-funded via MAICPP + SA so no budget compete.",
             "VBDs A2, A3 (lightweight)"),
            ("\"Our last digital initiative stalled.\"", "CEO, COO",
             "Stalls are pattern not luck \u2014 no measurement + no champions. A3 instruments adoption from day one; re-scope if metrics flat by week 6.",
             "VBDs A3, A8, D3"),
            ("\"Will AI take our jobs?\"", "LOB head, HR, Works Council",
             "AI changes job content, not headcount in wave one. Track hours-recovered-per-role; co-design reinvestment with HR. Never sequence rollout to a RIF.",
             "VBDs A2, D2"),
            ("\"We don't have AI skills internally.\"", "CIO, CHRO",
             "D2 maps roles to MS AI learning paths; MAICPP partners carry HoK while we transfer skills. Don't need full capability before starting.",
             "VBD D2 \u00b7 VBD B2"),
            ("\"Adoption will fragment by department.\"", "CIO, CDO",
             "Shadow AI already happening; A8 stands up your internal Copilot CoE with charter + single intake. Coordination beats ban-and-drive-underground.",
             "VBDs A8, B6"),
        ],
    },
    {
        "num": 3,
        "name": "AI Strategy & Experience",
        "subtitle": "Surface choice, capability scepticism, hallucinations",
        "narrative": (
            "Three surfaces, three jobs: M365 Copilot for productivity, "
            "Copilot Studio for processes, Foundry for code-first / "
            "customer-facing. Most CoEs run all three eventually. Lead with "
            "persona-to-surface mapping, not the SKU list."
        ),
        "owner": "Pillar 3 lead \u00b7 bring Foundry SME for technical depth",
        "objections": [
            ("\"Which surface \u2014 M365, Studio, or Foundry?\"", "CIO, CTO, CDO",
             "M365 for knowledge-worker productivity; Studio for low-code agents in processes; Foundry for code-first production. Horizon Assessment scores in a week.",
             "Horizon Assessment \u00b7 VBDs A1, B1, C9"),
            ("\"Hallucinations are a deal-breaker.\"", "CISO, Legal, regulated sponsor",
             "For ungrounded GenAI \u2014 yes. Production pattern is grounding + evals + human-in-loop (C13 + C15). Don't deploy unbounded chat to a regulated workflow.",
             "VBDs C13, C15"),
            ("\"We need agents, not just chat.\"", "CTO, Head of Innovation",
             "Agreed \u2014 chat is entry, agents are value. B1 first build, B2 co-build with partner, C18 (ATO Agent Factory) for industrialised delivery. 90 days to production.",
             "VBDs B1, B2, C18 \u00b7 A5"),
            ("\"GenAI is hype \u2014 we want classic ML.\"", "CDO, Chief Actuary",
             "You're right not every problem is GenAI. Foundry hosts both AzureML + Foundry with shared MLOps. Often the winning pattern is hybrid \u2014 ML predicts, GenAI explains.",
             "VBDs C19, C5"),
            ("\"Open-source models are better/cheaper.\"", "CTO, Head of Engineering",
             "Model choice matters. Foundry hosts AOAI + Mistral + Llama + DeepSeek + open weights behind one API with shared safety/eval. Migrated models in under a sprint.",
             "VBD C8 \u00b7 Foundry catalogue"),
            ("\"Show us something for OUR industry.\"", "LOB, Industry sponsor",
             "Eskom briefing is the regulated-industry template (POPIA, PFMA, NERSA); adapts to banking, public, healthcare. Or stand up A1 with YOUR scenarios in a week.",
             "Eskom briefing \u00b7 VBD A1"),
        ],
    },
    {
        "num": 4,
        "name": "Technology & Data",
        "subtitle": "Readiness, integration, lock-in, cost-overrun",
        "narrative": (
            "\"Our data isn't clean enough\" is the modal blocker. Counter: "
            "start where data is good and scope around what isn't. The "
            "Factory (F1\u2013F10) puts the foundation under everything at "
            "zero cost \u2014 customer-funded work is business logic only."
        ),
        "owner": "Pillar 4 lead \u00b7 bring FastTrack, FinOps SME",
        "objections": [
            ("\"Our data isn't clean enough.\"", "CDO, CIO",
             "True estate-wide; rarely true for one scoped use-case. Fabric SA (C7) finds AI-ready domains today; F3+F4 modernise lakehouse free in background.",
             "VBDs C7, C5 \u00b7 F3, F4"),
            ("\"[SAP/Mainframe] integration is the blocker.\"", "CIO, CTO",
             "Agentic AI changes unit cost. B2/B5 agents sit on top of SAP via existing API/RFC layer; abstract legacy for user without re-architecting SoR.",
             "VBDs B2, B5 \u00b7 F7"),
            ("\"We're worried about Azure lock-in.\"", "CTO, Procurement",
             "Lock-in worth worrying about is at data + model layer \u2014 not platform. Data stays in your tenant; Foundry is multi-model; agents are portable IP.",
             "VBD C8 \u00b7 Pillar 5 residency"),
            ("\"Landing zone isn't ready \u2014 18 months.\"", "CIO, Head of Cloud",
             "Factory math: F1+F2+F8+F9 zero-cost, MS-delivered under USS or 2-page agreement. AI-ready LZ in 6\u201310 weeks. Customer-funded work = business logic only.",
             "F1, F2, F8, F9 \u00b7 VBD C11"),
            ("\"We can't add another platform.\"", "CTO, Architect",
             "CoE asks you to consolidate on Azure + M365 you already have (~60\u201380% in place). F4/F5/F7/F10 net-reduce platform count by migrating off legacy.",
             "F4, F5, F7, F10"),
            ("\"Costs will spiral \u2014 we've seen Azure explode.\"", "CFO, FinOps",
             "WAF-for-AI (C12) bakes in token budgets + model routing + cache. D3 reports consumption-vs-value monthly. ACO/ECIF envelopes put overrun risk on Microsoft.",
             "VBDs C12, D3 \u00b7 ATO envelope"),
        ],
    },
    {
        "num": 5,
        "name": "Governance & Security",
        "subtitle": "POPIA, sovereignty, model risk, IP, prompt injection, EU AI Act",
        "narrative": (
            "In RSA this pillar is non-negotiable: POPIA, PFMA for SOEs, "
            "sector regulators (PA, FSCA, NERSA), and EU AI Act for cross-"
            "border firms. The Eskom Executive Briefing is the canonical "
            "regulated-industry response template."
        ),
        "owner": "Pillar 5 lead \u00b7 bring Legal, Compliance, Purview SMEs",
        "objections": [
            ("\"Data must stay in SA (POPIA / sovereignty).\"", "CISO, Legal, CIO, CDO",
             "Supported. SA North + West regions with paired residency. Copilot processes in-geo; Foundry data-zone-scoped. Eskom briefing walks POPIA/NERSA/PFMA slide-by-slide.",
             "Eskom briefing \u00b7 Purview AI Hub"),
            ("\"How do we manage model risk + explainability?\"", "CISO, CRO, Internal Audit",
             "RAI v2 (C15) gives Impact Assessment + fairness + reliability. GenAIOps (C13) gives continuous eval + red-team. Maps to NIST AI RMF + EU AI Act high-risk.",
             "VBDs C13, C15 \u00b7 RAI v2"),
            ("\"IP / copyright exposure on outputs?\"", "GC, CISO",
             "Microsoft's Copilot Copyright Commitment indemnifies customers vs 3rd-party IP claims on Copilot output (M365 Copilot, CfS, AOAI) when guardrails are configured.",
             "VBD A6 \u00b7 Copilot Copyright Commitment"),
            ("\"Shadow AI is already a problem.\"", "CISO, CIO",
             "Worst response is a ban \u2014 usage moves to phones. Give sanctioned alternative (Copilot Chat / Studio agent) + Purview AI Hub discovery. A6 produces report in 2 wks.",
             "VBD A6 \u00b7 Purview AI Hub"),
            ("\"Prompt injection / jailbreaks \u2014 defence?\"", "CISO, AppSec lead",
             "Three layers: Azure AI Content Safety + Prompt Shields at boundary; GenAIOps telemetry into Sentinel (F9 free); human-in-loop for irreversible actions.",
             "VBDs C9, C13 \u00b7 F9 Sentinel"),
            ("\"EU AI Act \u2014 are you compliant?\"", "CCO, Legal",
             "Maps to controls you'd put in place anyway. RAI v2 (C15) covers Art 9/13; GenAIOps (C13) covers Art 12/15. MS publishes per-product RAI Transparency Reports.",
             "VBDs C15, C13 \u00b7 RAI Transparency Reports"),
        ],
    },
    {
        "num": 6,
        "name": "Co-sell & Partner",
        "subtitle": "Incumbent SIs, partner overlap, hyperscaler choice",
        "narrative": (
            "Usually political, not technical. Don't take incumbents head-on "
            "\u2014 route work through them with MAICPP funding. Microsoft "
            "sells the platform; partners deliver; the CoE methodology + "
            "funding (ECIF, MAICPP, ACO) is the glue."
        ),
        "owner": "Pillar +1 lead + ATU partner manager",
        "objections": [
            ("\"We already work with [SI] \u2014 don't need MS to deliver.\"", "CIO, Procurement",
             "That IS the model. MS sells platform; MAICPP partners (often [their SI]) deliver. CoE provides methodology + ECIF/MAICPP/ACO funding + Factory underneath.",
             "Partner Recruitment Kit \u00b7 Op Playbook"),
            ("\"Why Microsoft and not [AWS / Google]?\"", "CTO, CIO",
             "Three AI-specific reasons: only MS has all 3 surfaces under one identity/data plane/RAI standard; $1M ATO envelope is MS-only; RSA in-country presence + 2 regions.",
             "Pitch Deck 3-surfaces \u00b7 ATO economics"),
            ("\"Partners add margin without value.\"", "CFO, Procurement",
             "MAICPP filters for AI-specialised partners with assessed competencies + references. Partner Scorecard shown to you. MAICPP co-funding offsets margin.",
             "Partner Scorecard \u00b7 MAICPP details"),
            ("\"Microsoft and partner tell us different things.\"", "CIO",
             "Delivery hygiene problem \u2014 we own it. RACI is explicit: MS owns MCEM 1\u20132, joint 3, partner-led 4\u20135 w/ MS assurance. Reset with RACI as contract.",
             "Delivery RACI \u00b7 Op Playbook"),
        ],
    },
]


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------

def build_title(prs):
    s = new_slide(prs)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, C_BG)
    add_rect(s, 0, Inches(5.6), SLIDE_W, Inches(1.9), C_DARK)
    add_text(s, Inches(0.7), Inches(0.7), Inches(10), Inches(0.4),
             "AI CENTER OF EXCELLENCE \u00b7 SELLER ENABLEMENT",
             size=12, bold=True, color=C_ACCENT)
    add_text(s, Inches(0.7), Inches(1.3), Inches(12), Inches(1.4),
             "Objection Handling",
             size=54, bold=True, color=C_INK)
    add_text(s, Inches(0.7), Inches(2.7), Inches(12), Inches(1.4),
             "Recurring customer pushback, mapped to the 5+1 Pillars,\nwith seller-ready responses and proof points in the pack.",
             size=22, color=C_MUTED)
    add_text(s, Inches(0.7), Inches(4.6), Inches(12), Inches(0.8),
             "Internal use \u00b7 ATU \u00b7 STU \u00b7 CSAM \u00b7 MAICPP partners",
             size=14, italic=True, color=C_HIGHLIGHT)
    add_text(s, Inches(0.7), Inches(5.85), Inches(12), Inches(0.4),
             "Yuri Baijnath", size=18, bold=True, color=C_BG)
    add_text(s, Inches(0.7), Inches(6.25), Inches(12), Inches(0.4),
             "CSU Cloud & AI Lead (South Africa)  \u00b7  Yuri.Baijnath@microsoft.com  \u00b7  Microsoft",
             size=12, color=C_ACCENT2)
    add_text(s, Inches(0.7), Inches(6.7), Inches(12), Inches(0.4),
             "Source: ai-coe-objection-handling-report.md  \u00b7  Re-verify pricing + ECIF caps before customer commit",
             size=10, italic=True, color=C_RULE)


def build_how_to_use(prs):
    s = new_slide(prs)
    add_section_tag(s, "How to use this deck")
    add_title(s, "Find the pillar, find the objection, deliver the response",
              "Every objection slide names the persona, the underlying concern, and the artefact to reach for")
    add_bullets(s, Inches(0.7), Inches(2.1), Inches(12), Inches(4), [
        "Address the concern, not the words \u2014 \"we'll wait\" is usually fear of being first; \"data isn't clean\" usually means \"give me a reason to defer.\"",
        "Match the response to the persona \u2014 a CFO objection takes a CFO answer (unit economics), not a technical one.",
        "Always close with the next artefact \u2014 Horizon Assessment, Pitch Deck slide, Eskom briefing, or a specific VBD.",
        "Top-10 battle card (slide 17) is the print-and-pin version for live calls.",
        "Escalation matrix (slide 18) names the CoE pillar lead to bring in when an objection blocks the deal.",
        "Re-verify any VBD funding or pricing on MCAPS Catalog before quoting it to a customer.",
    ], size=14)
    add_rect(s, Inches(0.6), Inches(6.55), Inches(12.13), Inches(0.45), C_BAND2)
    add_text(s, Inches(0.8), Inches(6.6), Inches(12), Inches(0.4),
             "Sequence for first meetings: lead with Pillar 1 (value) and Pillar 5 (governance); the others surface in meeting 2.",
             size=11, italic=True, color=C_MUTED)
    add_footer(s, 2)


def build_pillars_refresher(prs):
    s = new_slide(prs)
    add_section_tag(s, "Framework refresher")
    add_title(s, "The 5+1 Pillars \u2014 where objections live",
              "Every objection in this deck routes to one of these pillars and one owner")
    pillars_short = [
        ("1", "Business Strategy",   "Value, ROI, exit cost",       C_ACCENT),
        ("2", "Org & Culture",       "Adoption, skills, change",    C_HIGHLIGHT),
        ("3", "AI Strategy & XP",    "Surface choice, hallucinations", C_DARK),
        ("4", "Tech & Data",         "Readiness, lock-in, cost",    C_GOOD),
        ("5", "Governance & Sec",    "POPIA, model risk, IP",       C_WARN),
        ("+1", "Co-sell & Partner",  "SI, hyperscaler, MAICPP",     C_ACCENT2),
    ]
    px = Inches(0.6); py = Inches(2.2); pw = Inches(2.02); ph = Inches(1.7)
    for n, name, sub, col in pillars_short:
        add_rect(s, px, py, pw, ph, C_BAND, line=C_RULE)
        add_rect(s, px, py, pw, Inches(0.06), col)
        add_text(s, px + Inches(0.12), py + Inches(0.15), pw - Inches(0.2), Inches(0.35),
                 f"Pillar {n}", size=10, bold=True, color=col)
        add_text(s, px + Inches(0.12), py + Inches(0.48), pw - Inches(0.2), Inches(0.5),
                 name, size=13, bold=True, color=C_INK)
        add_text(s, px + Inches(0.12), py + Inches(1.0), pw - Inches(0.2), Inches(0.7),
                 sub, size=11, color=C_MUTED)
        px += pw + Inches(0.01)

    add_rect(s, Inches(0.6), Inches(4.2), Inches(12.13), Inches(2.6), C_DARK)
    add_text(s, Inches(0.85), Inches(4.35), Inches(11.5), Inches(0.4),
             "WHY OBJECTIONS MAP TO PILLARS", size=11, bold=True, color=C_ACCENT2)
    add_bullets(s, Inches(0.85), Inches(4.7), Inches(11.5), Inches(2.0), [
        "Routing \u2014 the right CoE owner takes the response, not whoever happens to be in the room.",
        "Coverage \u2014 you can see immediately if a customer's objections cluster on one pillar (often Pillar 5 for regulated, Pillar 1 for cost-pressed).",
        "Proof points \u2014 each pillar already has artefacts in the pack; you reach for the right one without improvising.",
        "Escalation \u2014 if an objection blocks the deal, the pillar lead is the named escalation path.",
    ], size=12, color=C_BG, bullet_color=C_ACCENT2)
    add_footer(s, 3)


def build_objection_map(prs):
    s = new_slide(prs)
    add_section_tag(s, "Objection map")
    add_title(s, "Where objections cluster by persona",
              "Read down to find your persona; read across to see the pillars they pressure")
    headers = ["Persona", "P1 Biz", "P2 Org", "P3 AI XP", "P4 Tech", "P5 Gov", "+1 Partner"]
    rows = [
        ["CEO",            "ROI \u00b7 Wait", "Stall risk",  "\u2014",          "\u2014",            "\u2014",            "\u2014"],
        ["CFO",            "ROI \u00b7 Exit \u00b7 Per-seat", "\u2014", "\u2014", "Costs spiral",     "\u2014",            "Partner margin"],
        ["CIO",            "Wait \u00b7 Build vs buy", "Skills \u00b7 Fragment", "Surface choice", "Integration \u00b7 LZ \u00b7 Lock-in", "POPIA", "Why MS \u00b7 SI conflict"],
        ["CISO / CRO",     "\u2014",         "\u2014",       "Hallucinations",  "\u2014",            "All",               "\u2014"],
        ["CDO",            "\u2014",         "\u2014",       "GenAI vs ML",     "Data quality",      "POPIA",             "\u2014"],
        ["CTO",            "Build vs buy",   "\u2014",       "Open source \u00b7 Agents", "Lock-in \u00b7 Tech debt", "\u2014", "Why MS"],
        ["CHRO / COO",     "\u2014",         "Readiness \u00b7 Fatigue \u00b7 Jobs", "\u2014", "\u2014", "\u2014",  "\u2014"],
        ["LOB head",       "\u2014",         "Fatigue \u00b7 Jobs", "Industry-specific", "\u2014",    "\u2014",            "\u2014"],
        ["GC / Legal",     "\u2014",         "\u2014",       "Hallucinations",  "\u2014",            "IP \u00b7 EU AI Act", "\u2014"],
        ["Procurement",    "Exit cost",      "\u2014",       "\u2014",          "Lock-in",           "\u2014",            "Partner \u00b7 Why MS"],
    ]
    add_table(s, Inches(0.6), Inches(2.05), Inches(12.13), Inches(4.7),
              headers, rows, body_size=10, header_size=11,
              col_widths=[1.6, 1.95, 1.95, 1.95, 1.95, 1.6, 1.95])
    add_text(s, Inches(0.6), Inches(6.85), Inches(12), Inches(0.2),
             "Diagonal density tells you where the customer's pain sits. If 70% lands on one pillar, escalate to that lead before meeting 2.",
             size=10, italic=True, color=C_MUTED)
    add_footer(s, 4)


def build_pillar_overview(prs, pillar, slide_no):
    """Overview slide for a single pillar."""
    s = new_slide(prs)
    n = pillar["num"]
    name = pillar["name"]
    col = PILLAR_COLORS[n]
    label = "+1" if n == 6 else str(n)

    add_section_tag(s, f"Pillar {label} \u00b7 {name}", color=col)
    add_title(s, f"Pillar {label} \u2014 {name}", pillar["subtitle"])

    # Left card: narrative
    add_rect(s, Inches(0.6), Inches(2.1), Inches(7.7), Inches(4.6), C_BG, line=col)
    add_rect(s, Inches(0.6), Inches(2.1), Inches(7.7), Inches(0.06), col)
    add_text(s, Inches(0.8), Inches(2.25), Inches(7.3), Inches(0.4),
             "HOW TO POSITION THIS PILLAR", size=11, bold=True, color=col)
    add_text(s, Inches(0.8), Inches(2.6), Inches(7.3), Inches(2.0),
             pillar["narrative"], size=15, color=C_INK)
    add_text(s, Inches(0.8), Inches(5.7), Inches(7.3), Inches(0.4),
             "ESCALATION OWNER", size=11, bold=True, color=col)
    add_text(s, Inches(0.8), Inches(6.05), Inches(7.3), Inches(0.6),
             pillar["owner"], size=13, color=C_INK)

    # Right card: objection count + persona summary
    add_rect(s, Inches(8.45), Inches(2.1), Inches(4.28), Inches(4.6), C_DARK)
    add_text(s, Inches(8.6), Inches(2.25), Inches(4.0), Inches(0.4),
             "OBJECTIONS IN THIS PILLAR", size=11, bold=True, color=C_ACCENT2)
    add_text(s, Inches(8.6), Inches(2.65), Inches(4.0), Inches(1.2),
             str(len(pillar["objections"])),
             size=80, bold=True, color=C_BG)
    add_text(s, Inches(8.6), Inches(4.05), Inches(4.0), Inches(0.4),
             "TYPICAL PERSONAS", size=11, bold=True, color=C_ACCENT2)
    # Collect unique personas
    personas = []
    for _, p, _, _ in pillar["objections"]:
        for token in [t.strip() for t in p.split(",")]:
            if token and token not in personas:
                personas.append(token)
    add_text(s, Inches(8.6), Inches(4.4), Inches(4.0), Inches(2.2),
             "\n".join(f"\u2022 {p}" for p in personas[:8]),
             size=12, color=C_BG)
    add_footer(s, slide_no)


def build_pillar_table(prs, pillar, slide_no):
    """Objection menu table for a single pillar."""
    s = new_slide(prs)
    n = pillar["num"]
    name = pillar["name"]
    col = PILLAR_COLORS[n]
    label = "+1" if n == 6 else str(n)

    add_section_tag(s, f"Pillar {label} \u00b7 {name} \u00b7 objection menu", color=col)
    add_title(s, f"Recurring objections \u2014 Pillar {label}",
              "Short response + proof point. Long-form rationale in ai-coe-objection-handling-report.md.")

    rows = [[obj, persona, resp, proof]
            for (obj, persona, resp, proof) in pillar["objections"]]
    add_table(s, Inches(0.6), Inches(2.1), Inches(12.13), Inches(4.7),
              ["Objection", "Persona", "Response (short)", "Proof point"],
              rows, body_size=10, header_size=11, header_fill=col,
              col_widths=[2.6, 1.9, 5.4, 2.23])
    add_text(s, Inches(0.6), Inches(6.85), Inches(12), Inches(0.2),
             f"Owner if it blocks the deal: {pillar['owner']}",
             size=10, italic=True, color=C_MUTED)
    add_footer(s, slide_no)


def build_battle_card(prs, slide_no):
    s = new_slide(prs)
    add_section_tag(s, "Top-10 battle card")
    add_title(s, "Print and pin \u2014 the ten you will hear most often",
              "One-line response per objection. Long form on the pillar slides; full rationale in the report.")
    rows = [
        ["1",  "\"ROI is unproven.\"",                "CoE forces value case per use-case; ATO absorbs first $1M of envelope; D3 reports CFO-grade outcomes quarterly."],
        ["2",  "\"We'll wait.\"",                     "88% of orgs already in (McKinsey 2025); T1 envisioning is free \u2014 stop after if value case doesn't hold."],
        ["3",  "\"Copilot per-seat doesn't pencil.\"", "Per-seat \u00d7 headcount is the wrong denominator; A2 produces hours-recovered-per-persona before purchase."],
        ["4",  "\"We'll build it ourselves.\"",       "Building is cheap; operating safely at scale is the stall. Co-build keeps IP yours, run-time discipline ours."],
        ["5",  "\"Data isn't clean enough.\"",        "Start where it is; Fabric SA scopes that. Worst data is the data no one reads."],
        ["6",  "\"Hallucinations are a deal-breaker.\"", "For ungrounded \u2014 yes. Production pattern: grounding + evals + human-in-loop (C13 + C15)."],
        ["7",  "\"Which surface?\"",                  "M365 for productivity; Studio for processes; Foundry for code-first. Horizon Assessment scores in a week."],
        ["8",  "\"Data must stay in SA.\"",           "SA North + West regions; Copilot processes in-geo; Foundry data-zone-scoped. Eskom briefing covers POPIA."],
        ["9",  "\"Shadow AI is a problem.\"",         "Don't ban \u2014 sanctioned alternative + Purview AI Hub. A6 produces discovery report in 2 weeks."],
        ["10", "\"We already have an SI.\"",          "That IS the model \u2014 they deliver, we fund (MAICPP / ECIF / ACO) and bring the methodology."],
    ]
    add_table(s, Inches(0.6), Inches(2.05), Inches(12.13), Inches(4.85),
              ["#", "Objection", "One-line response"],
              rows, body_size=11, header_size=11,
              col_widths=[0.5, 3.5, 8.13])
    add_footer(s, slide_no)


def build_escalation_matrix(prs, slide_no):
    s = new_slide(prs)
    add_section_tag(s, "Escalation matrix")
    add_title(s, "When an objection blocks the deal \u2014 who do you bring in?",
              "Named CoE owner per pillar + the SME to bring with them")
    rows = [
        ["1 \u00b7 Business Strategy",     "\"ROI is unproven\" \u00b7 \"Exit cost?\"",          "Pillar 1 lead",           "CSU GM, ATU lead"],
        ["2 \u00b7 Org & Culture",         "\"Adoption will stall\" \u00b7 \"AI takes jobs\"",   "Pillar 2 lead",           "Customer HR sponsor"],
        ["3 \u00b7 AI Strategy & XP",      "\"Hallucinations\" \u00b7 \"Which surface?\"",       "Pillar 3 lead",           "Foundry SME"],
        ["4 \u00b7 Tech & Data",           "\"Data isn't clean\" \u00b7 \"Costs will spiral\"",  "Pillar 4 lead",           "FastTrack, FinOps SME"],
        ["5 \u00b7 Governance & Security", "\"Data residency\" \u00b7 \"Model risk\"",           "Pillar 5 lead",           "Legal, Compliance, Purview SME"],
        ["+1 \u00b7 Co-sell & Partner",    "\"MS vs partner conflict\"",                          "ATU partner manager",     "MAICPP PSE"],
    ]
    add_table(s, Inches(0.6), Inches(2.1), Inches(12.13), Inches(4.0),
              ["Pillar", "Hardest objection", "Owner (CoE side)", "Bring also"],
              rows, body_size=11, header_size=11,
              col_widths=[2.7, 4.0, 2.5, 2.93])
    add_rect(s, Inches(0.6), Inches(6.3), Inches(12.13), Inches(0.65), C_BAND)
    add_text(s, Inches(0.8), Inches(6.4), Inches(12), Inches(0.5),
             "Rule of thumb: escalate after the same objection lands twice from different stakeholders. Once is friction; twice is a pattern.",
             size=12, italic=True, color=C_INK)
    add_footer(s, slide_no)


def build_linked_artefacts(prs, slide_no):
    s = new_slide(prs)
    add_section_tag(s, "Linked artefacts")
    add_title(s, "What to reach for, by objection theme",
              "Don't improvise \u2014 every theme already has an artefact in the pack")
    rows = [
        ["ROI / value scepticism",       "AI-CoE-Pitch-Deck.pptx (ATO slide) \u00b7 VBD D3 Value Realization"],
        ["Maturity / where to start",    "AI-CoE-Horizon-Assessment.xlsx"],
        ["Regulated-industry concerns",  "AI-CoE-Eskom-Executive-Briefing.pptx"],
        ["C-suite leave-behind needed",  "AI-CoE-Executive-OnePager.docx"],
        ["Partner / SI politics",        "AI-CoE-Partner-Recruitment-Kit.pptx \u00b7 AI-CoE-Partner-Scorecard.xlsx"],
        ["Delivery governance challenge","AI-CoE-Operating-Playbook.docx \u00b7 AI-CoE-Delivery-RACI.xlsx"],
        ["VBD detail requested",         "AI-CoE-VBD-Reference-Deck.pptx \u00b7 csu-ai-vbd-reference-report.md"],
        ["Seller onboarding to the pack","AI-CoE-How-To-Use.pptx"],
    ]
    add_table(s, Inches(0.6), Inches(2.1), Inches(12.13), Inches(4.5),
              ["When this objection lands", "Reach for"],
              rows, body_size=12, header_size=11,
              col_widths=[4.0, 8.13])
    add_text(s, Inches(0.6), Inches(6.7), Inches(12), Inches(0.25),
             "Every link is a file in the root of the AI CoE repo. If it's not there, the artefact hasn't been built yet \u2014 raise via improvements-plan.md.",
             size=10, italic=True, color=C_MUTED)
    add_footer(s, slide_no)


def build_sources(prs, slide_no):
    s = new_slide(prs)
    add_section_tag(s, "Sources & methodology")
    add_title(s, "Where these objections and responses came from",
              "Honest disclosure \u2014 re-verify any pricing or named GA dates before customer commit")
    add_bullets(s, Inches(0.7), Inches(2.1), Inches(12), Inches(4.4), [
        "HIGH trust \u2014 ai-coe-objection-handling-report.md (Yuri Baijnath, author).",
        "HIGH trust \u2014 ATO funding envelope numbers cross-checked with internal ATO deck.",
        "HIGH trust \u2014 csu-ai-vbd-reference-report.md (VBD references throughout).",
        "MODERATE trust \u2014 BCG 2024 (\"60% reap little value\"; \"1/3 trained\"), McKinsey 2025 State of AI (\"88% using\"). Verified in ai-coe-pitch-stats-report.md.",
        "Pillar definitions \u2014 internal Frontier CoE pillar decks (5+1).",
        "Methodology \u2014 objections seeded from RSA CSU recurring patterns; responses drafted against existing VBDs + Pitch-Deck proof points; battle-card distilled to one-liners that survive interruption.",
        "Re-verify \u2014 ECIF caps, ATO envelope size, regional residency policy, and any quoted VBD funding before commit (MCAPS Catalog).",
    ], size=12, line_spacing=1.25)
    add_rect(s, Inches(0.6), Inches(6.6), Inches(12.13), Inches(0.4), C_BAND2)
    add_text(s, Inches(0.7), Inches(6.62), Inches(12), Inches(0.35),
             "Limitations: objections are RSA-CSU pattern-derived, not exhaustive. New objections should be raised back into ai-coe-objection-handling-report.md.",
             size=10, italic=True, color=C_MUTED)
    add_footer(s, slide_no)


def build_close(prs):
    s = new_slide(prs)
    add_rect(s, 0, 0, SLIDE_W, SLIDE_H, C_DARK)
    add_text(s, Inches(0.7), Inches(0.7), Inches(12), Inches(0.4),
             "AI CENTER OF EXCELLENCE  \u00b7  CLOSING",
             size=12, bold=True, color=C_ACCENT2)
    add_text(s, Inches(0.7), Inches(1.3), Inches(12), Inches(1.4),
             "Address the concern, not the words.",
             size=48, bold=True, color=C_BG)
    add_text(s, Inches(0.7), Inches(2.7), Inches(12), Inches(0.7),
             "Every objection in this deck routes to a pillar, a response,\nand the next artefact to put on the table.",
             size=22, color=C_ACCENT2)
    nums = [
        ("5+1",  "Pillars covered"),
        ("34",   "Objections catalogued"),
        ("10",   "On the battle card"),
        ("1",    "Source-of-truth report"),
    ]
    nx = Inches(0.7)
    for n, lbl in nums:
        add_text(s, nx, Inches(4.4), Inches(3.0), Inches(0.7),
                 n, size=44, bold=True, color=C_HIGHLIGHT)
        add_text(s, nx, Inches(5.1), Inches(3.0), Inches(0.4),
                 lbl, size=12, color=C_ACCENT2)
        nx += Inches(3.0)
    add_rect(s, Inches(0.7), Inches(6.05), Inches(11.9), Emu(9000), C_ACCENT2)
    add_text(s, Inches(0.7), Inches(6.2), Inches(12), Inches(0.4),
             "Yuri Baijnath", size=18, bold=True, color=C_BG)
    add_text(s, Inches(0.7), Inches(6.6), Inches(12), Inches(0.4),
             "CSU Cloud & AI Lead (South Africa)  \u00b7  Yuri.Baijnath@microsoft.com  \u00b7  Microsoft",
             size=12, color=C_ACCENT2)


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------

def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    build_title(prs)                             # 1
    build_how_to_use(prs)                        # 2
    build_pillars_refresher(prs)                 # 3
    build_objection_map(prs)                     # 4
    # 6 pillars x 2 slides each = 12 slides (5..16)
    slide_no = 5
    for pillar in PILLARS:
        build_pillar_overview(prs, pillar, slide_no)
        slide_no += 1
        build_pillar_table(prs, pillar, slide_no)
        slide_no += 1
    build_battle_card(prs, slide_no);          slide_no += 1  # 17
    build_escalation_matrix(prs, slide_no);    slide_no += 1  # 18
    build_linked_artefacts(prs, slide_no);     slide_no += 1  # 19
    build_sources(prs, slide_no);              slide_no += 1  # 20
    build_close(prs)                                              # 21

    out = "AI-CoE-Objection-Handling.pptx"
    prs.save(out)
    print(f"Wrote {out} with {len(prs.slides)} slides.")


if __name__ == "__main__":
    main()
